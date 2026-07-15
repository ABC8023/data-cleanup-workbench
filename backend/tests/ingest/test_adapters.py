from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from data_workbench.engine.duckdb_runtime import DuckDBRuntime
from data_workbench.ingest.base import MalformedInput, UnsupportedFormat
from data_workbench.ingest.registry import AdapterRegistry


@pytest.fixture
def fixture_dir(tmp_path: Path) -> Path:
    fixtures = tmp_path / "fixtures"
    fixtures.mkdir()

    (fixtures / "simple.csv").write_text("id,name\n1,Ada\n", encoding="utf-8")
    (fixtures / "records.json").write_text(
        json.dumps([{"id": 1, "profile": {"name": "Ada"}, "tags": ["admin"]}]),
        encoding="utf-8",
    )
    (fixtures / "records.ndjson").write_text(
        json.dumps({"id": 1, "profile": {"name": "Lin"}}) + "\n",
        encoding="utf-8",
    )
    pq.write_table(pa.table({"id": [1], "name": ["Grace"]}), fixtures / "simple.parquet")

    workbook = openpyxl.Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.append(["id", "amount"])
    orders.append([1, 12.5])
    customers = workbook.create_sheet("Customers")
    customers.append(["id", "name"])
    customers.append([1, "Ada"])
    workbook.save(fixtures / "book.xlsx")

    return fixtures


@pytest.mark.parametrize(
    ("fixture", "expected"),
    [
        ("simple.csv", ["data"]),
        ("records.json", ["data"]),
        ("records.ndjson", ["data"]),
        ("simple.parquet", ["data"]),
        ("book.xlsx", ["Orders", "Customers"]),
    ],
)
def test_registry_returns_queryable_tables(
    fixture: str,
    expected: list[str],
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    handles = AdapterRegistry().inspect(fixture_dir / fixture, tmp_path)

    assert [handle.name for handle in handles] == expected
    with DuckDBRuntime("512MB", 1).connect(tmp_path) as connection:
        assert all(
            connection.sql(f"SELECT count(*) FROM ({handle.scan_sql})").fetchone()[0] > 0
            for handle in handles
        )


@pytest.fixture
def invalid_fixture_dir(tmp_path: Path) -> Path:
    fixtures = tmp_path / "invalid"
    fixtures.mkdir()
    (fixtures / "corrupt.csv").write_bytes(b"\x00\xff\x00\xfe")
    (fixtures / "ambiguous.json").write_text('{"id": 1}', encoding="utf-8")

    empty = openpyxl.Workbook()
    empty.save(fixtures / "empty.xlsx")

    duplicate_sheets = openpyxl.Workbook()
    duplicate_sheets.active.title = "Data"
    duplicate_sheets.active.append(["id"])
    duplicate_sheets.active.append([1])
    duplicate_sheets.create_sheet(" data ")
    duplicate_sheets.save(fixtures / "duplicate-sheets.xlsx")

    (fixtures / "legacy.xls").write_bytes(b"legacy")
    (fixtures / "renamed.exe.csv").write_bytes(b"MZ\x90\x00program")
    return fixtures


@pytest.mark.parametrize(
    "name",
    [
        "corrupt.csv",
        "ambiguous.json",
        "empty.xlsx",
        "duplicate-sheets.xlsx",
        "legacy.xls",
        "renamed.exe.csv",
    ],
)
def test_invalid_inputs_raise_typed_errors(
    name: str,
    invalid_fixture_dir: Path,
    tmp_path: Path,
) -> None:
    with pytest.raises((MalformedInput, UnsupportedFormat)):
        AdapterRegistry().inspect(invalid_fixture_dir / name, tmp_path)


def test_parquet_extension_requires_parquet_bytes(tmp_path: Path) -> None:
    source = tmp_path / "renamed.parquet"
    source.write_text("id,name\n1,Ada\n", encoding="utf-8")

    with pytest.raises(MalformedInput, match="Parquet"):
        AdapterRegistry().inspect(source, tmp_path / "session")


def test_json_staging_is_deterministic_and_preserves_source(tmp_path: Path) -> None:
    source = tmp_path / "source's.json"
    source.write_text(
        json.dumps(
            [
                {"z": 2, "meta": {"name": "Ada"}, "values": [2.5, {"b": 2, "a": 1}]},
                {"a": True, "meta": {"name": "Lin"}},
            ]
        ),
        encoding="utf-8",
    )
    original = source.read_bytes()
    session_dir = tmp_path / "session"

    handle = AdapterRegistry().inspect(source, session_dir)[0]

    assert source.read_bytes() == original
    assert (session_dir / "staged-json.parquet").is_file()
    with DuckDBRuntime("512MB", 1).connect(session_dir) as connection:
        relation = connection.sql(handle.scan_sql)
        result = relation.fetchall()
        assert relation.columns == [
            "a",
            "meta.name",
            "values",
            "z",
        ]
    assert result == [
        (None, "Ada", '[2.5,{"a":1,"b":2}]', "2"),
        ("true", "Lin", None, None),
    ]


def test_duckdb_runtime_applies_resource_and_extension_policy(tmp_path: Path) -> None:
    session_dir = tmp_path / "session"

    with DuckDBRuntime("512MB", 2).connect(session_dir) as connection:
        settings = connection.execute(
            """
            SELECT current_setting('memory_limit'),
                   current_setting('threads'),
                   current_setting('temp_directory'),
                   current_setting('allow_unsigned_extensions')
            """
        ).fetchone()

    assert settings is not None
    assert settings[0].startswith("488")
    assert settings[1] == 2
    assert Path(settings[2]) == session_dir / "duckdb.tmp"
    assert settings[3] is False
    assert DuckDBRuntime().memory_limit == "4GB"


@pytest.mark.parametrize("header", [[None, "name"], ["id", " ID "]])
def test_excel_rejects_invalid_normalized_headers(
    header: list[str | None],
    tmp_path: Path,
) -> None:
    source = tmp_path / "invalid-headers.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(header)
    workbook.active.append([1, "Ada"])
    workbook.save(source)

    with pytest.raises(MalformedInput, match="headers"):
        AdapterRegistry().inspect(source, tmp_path / "session")


def test_json_rejects_non_object_records_and_excess_columns(tmp_path: Path) -> None:
    session_dir = tmp_path / "session"
    non_object = tmp_path / "non-object.ndjson"
    non_object.write_text('[1, 2, 3]\n', encoding="utf-8")
    too_wide = tmp_path / "too-wide.json"
    too_wide.write_text(
        json.dumps([{f"column_{index}": index for index in range(10_001)}]),
        encoding="utf-8",
    )

    with pytest.raises(MalformedInput, match="objects"):
        AdapterRegistry().inspect(non_object, session_dir)
    with pytest.raises(MalformedInput, match="10000 columns"):
        AdapterRegistry().inspect(too_wide, session_dir)
